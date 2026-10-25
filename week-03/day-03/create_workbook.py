import csv
from pathlib import Path
from collections import defaultdict
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(__file__).resolve().parents[2]
source = root / "week-01" / "day-05" / "sales_data.csv"
out = root / "week-03" / "day-03" / "analysis_workbook.xlsx"

with source.open("r", newline="", encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    records = list(reader)

if not records:
    raise SystemExit("ERROR: sales_data.csv has no rows.")

def number(value):
    try:
        return float(value) if value not in (None, "") else 0.0
    except ValueError:
        return 0.0

def parse_date(value):
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except (ValueError, TypeError):
            pass
    return None

wb = Workbook()
ws = wb.active
ws.title = "Sales Data"
headers = list(records[0].keys())
ws.append(headers + ["Quarter", "Revenue", "Estimated Margin %", "Previous Row Growth %"])

previous_revenue = None
for rec in records:
    row = [rec.get(h, "") for h in headers]
    amount = number(rec.get("amount"))
    date = parse_date(rec.get("sale_date", ""))
    quarter = f"Q{((date.month - 1) // 3) + 1} {date.year}" if date else "Unknown"
    discount = number(rec.get("discount"))
    # Margin is an illustrative proxy only: discount is not profit/cost.
    estimated_margin = max(0, 100 - discount) / 100 if rec.get("discount") not in (None, "") else None
    growth = (amount / previous_revenue - 1) if previous_revenue not in (None, 0) else None
    ws.append(row + [quarter, amount, estimated_margin, growth])
    previous_revenue = amount

for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
for col in ws.columns:
    letter = col[0].column_letter
    ws.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 28)

data_table = Table(displayName="SalesTable", ref=ws.dimensions)
data_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
ws.add_table(data_table)

summary = wb.create_sheet("Region Summary")
summary.append(["Region", "Revenue"])
region_totals = defaultdict(float)
for rec in records:
    region_totals[rec.get("region", "Unknown")] += number(rec.get("amount"))
for region, total in sorted(region_totals.items()):
    summary.append([region, total])

for cell in summary[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
for row in range(2, summary.max_row + 1):
    summary.cell(row, 2).number_format = '#,##0.00'
summary.column_dimensions["A"].width = 20
summary.column_dimensions["B"].width = 18

chart = BarChart()
chart.title = "Revenue by Region"
chart.y_axis.title = "Revenue"
chart.x_axis.title = "Region"
chart.add_data(Reference(summary, min_col=2, min_row=1, max_row=summary.max_row), titles_from_data=True)
chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=summary.max_row))
summary.add_chart(chart, "D2")

for row in range(2, ws.max_row + 1):
    ws.cell(row, headers.index("amount") + 1).number_format = '#,##0.00'
    ws.cell(row, len(headers) + 3).number_format = '0.0%'
    ws.cell(row, len(headers) + 4).number_format = '0.0%'
ws.conditional_formatting.add(
    f"{ws.cell(2, len(headers)+4).coordinate}:{ws.cell(ws.max_row, len(headers)+4).coordinate}",
    CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="C6EFCE"))
)

notes = wb.create_sheet("Instructions")
notes.append(["Feature", "Details"])
notes.append(["Source", "Imported from week-01/day-05/sales_data.csv"])
notes.append(["Revenue", "Uses the amount column as revenue."])
notes.append(["Quarter", "Derived from sale_date when a date can be parsed."])
notes.append(["Growth %", "Row-to-row revenue change, not year-over-year growth."])
notes.append(["Margin %", "Illustrative placeholder based on discount; not actual profit margin. Actual margin requires profit and revenue/cost data."])
notes.append(["PivotTable / Power Query", "Use Excel Data > Get Data > From Text/CSV to import additional CSVs, then Insert > PivotTable. This workbook contains a source table and a regional summary chart but does not create a native PivotTable or Power Query query."])
for cell in notes[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
notes.column_dimensions["A"].width = 28
notes.column_dimensions["B"].width = 105
for row in notes.iter_rows():
    for cell in row:
        cell.alignment = cell.alignment.copy(wrap_text=True)

wb.save(out)
print(f"SUCCESS: Created {out}")
print(f"Rows imported: {len(records)}")
print(f"Regions summarized: {len(region_totals)}")
